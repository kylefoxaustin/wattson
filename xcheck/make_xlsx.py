#!/usr/bin/env python3
"""Regenerate RESULTS-50.xlsx from the CSVs. Committed so the workbook can be
re-derived and audited rather than taken on trust."""
import csv, os, statistics as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

G = PatternFill('solid', fgColor='C6EFCE')
A = PatternFill('solid', fgColor='FFEB9C')
R = PatternFill('solid', fgColor='FFC7CE')
HD = PatternFill('solid', fgColor='1F3864')
HF = Font(bold=True, color='FFFFFF')
NOTE = ('rails MEASURED on IMX95LPD5EVK-19 via NXP BCU (busy-window median, see busywin.py); '
        'A55 die 40-41 C under load; activity DERIVED from QEMU TCG plugins')

def shade(c, v):
    c.fill = G if v <= 8 else (A if v <= 12 else R)
    c.number_format = '0.0'

def header(ws, hdr):
    ws.append(hdr)
    for c in ws[1]:
        c.fill = HD; c.font = HF
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook(); wb.remove(wb.active)

# ── activity factors (Part 1) ─────────────────────────────────────────────
act = list(csv.DictReader(open('RESULTS-activity.csv')))
ws = wb.create_sheet('activity factors')
header(ws, ['app', 'QEMU insns', 'silicon insns', 'insn err %',
            'QEMU read GB/s', 'silicon read GB/s', 'read err %',
            'QEMU total GB/s (model input)', 'elapsed s'])
for r in act:
    be = float(r['read_err']) if r['read_err'] else None
    ws.append([r['app'], float(r['qemu_insns']), float(r['silicon_insns']), float(r['insn_err']),
               float(r['qemu_read_GBps']), float(r['silicon_read_GBps']), be,
               float(r['qemu_total_GBps']), float(r['elapsed_s'])])
    c = ws.cell(ws.max_row, 4); v = abs(float(r['insn_err']))
    c.fill = G if v <= 2 else (A if v <= 8 else R); c.number_format = '0.00'
    if be is not None and float(r['silicon_read_GBps']) >= 0.25:
        c2 = ws.cell(ws.max_row, 7)
        c2.fill = R if abs(be) > 60 else (A if abs(be) > 25 else G); c2.number_format = '0.0'
    for col in (2, 3): ws.cell(ws.max_row, col).number_format = '#,##0'
    for col in (5, 6, 8): ws.cell(ws.max_row, col).number_format = '0.000'
    ws.cell(ws.max_row, 1).font = Font(bold=True)
ie = [abs(float(r['insn_err'])) for r in act]
ws.append([])
ws.append([f'instructions: MAPE {st.mean(ie):.2f}%, median {st.median(ie):.2f}%, '
           f'{sum(1 for x in ie if x <= 2)} of {len(ie)} within 2%'])
ws.cell(ws.max_row, 1).font = Font(bold=True, color='1F3864')
ws.append(['l3d_cache_refill is a REFILL counter and sees READS; the like-for-like QEMU quantity is '
           'dram_read_proxy. The model is driven by TOTAL traffic (last column), which is the quantity '
           'its coefficient was calibrated against.'])
ws.cell(ws.max_row, 1).font = Font(italic=True, size=9, color='808080')
ws.append(['QEMU counts DERIVED from TCG plugins (linux-user); silicon counts MEASURED via ARM PMU '
           '(i.MX95, A55 core 0, pinned 1.8 GHz)'])
ws.cell(ws.max_row, 1).font = Font(italic=True, size=9, color='808080')
for col, w in zip('ABCDEFGHI', (14, 15, 15, 11, 14, 15, 11, 24, 10)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B2'

# ── single-app ────────────────────────────────────────────────────────────
rows = list(csv.DictReader(open('RESULTS-50.csv')))
ws = wb.create_sheet('50 apps (single)')
header(ws, ['app', 'ALU Mops/s', 'GB/s', 'pred arm', 'meas arm', 'err arm %',
            'pred soc', 'meas soc', 'err soc %', 'pred SUM', 'meas SUM', 'err SUM %',
            'iters', 'elapsed s', 'window'])
for r in rows:
    ws.append([r['app'], float(r['aluMps']), float(r['GBps']),
               float(r['pred_arm']), float(r['meas_arm']), float(r['err_arm']),
               float(r['pred_soc']), float(r['meas_soc']), float(r['err_soc']),
               float(r['pred_sum']), float(r['meas_sum']), float(r['err_sum']),
               int(r['iters']), float(r['elapsed_s']),
               f"[{r['win_t0']},{r['win_t1']}]s"])
    for col, key in ((6, 'err_arm'), (9, 'err_soc'), (12, 'err_sum')):
        shade(ws.cell(ws.max_row, col), float(r[key]))
    ws.cell(ws.max_row, 1).font = Font(bold=True)
for col, w in zip('ABCDEFGHIJKLMNO', (14, 11, 8, 9, 9, 9, 9, 9, 9, 9, 9, 10, 7, 9, 11)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B2'

# ── multi-app ─────────────────────────────────────────────────────────────
mc = list(csv.DictReader(open('RESULTS-multicore.csv')))
ws = wb.create_sheet('multi-app (concurrent)')
header(ws, ['what it is', 'applications running concurrently', 'cores', 'cores busy',
            'QEMU ALU Mops/s', 'QEMU GB/s', 'PMU ALU Mops/s', 'PMU GB/s',
            'pred (QEMU) mW', 'pred (PMU) mW', 'meas mW', 'err %', 'err PMU %'])
def emit(sel, label):
    ws.append([label]); ws.cell(ws.max_row, 1).font = Font(bold=True, italic=True)
    for r in sel:
        ws.append([r['kind'], r['apps'], int(r['cores']), float(r['cores_eff']),
                   float(r['qemu_aluMps']), float(r['qemu_GBps']),
                   float(r['pmu_aluMps']), float(r['pmu_GBps']),
                   float(r['pred_qemu']), float(r['pred_pmu']), float(r['meas_sum']),
                   float(r['err_qemu']), float(r['err_pmu'])])
        shade(ws.cell(ws.max_row, 12), float(r['err_qemu']))
        shade(ws.cell(ws.max_row, 13), float(r['err_pmu']))
        ws.cell(ws.max_row, 1).font = Font(bold=True)
edge = [r for r in mc if r['case'].startswith('e')]
ctrl = [r for r in mc if not r['case'].startswith('e')]
emit(edge, 'REALISTIC EDGE WORKLOADS')
ws.append([])
emit(ctrl, 'CONTROLS - synthetic / benchmark mixes')
ws.append([])
eq = [float(r['err_qemu']) for r in edge]
ws.append([f'edge mixes: mean {st.mean(eq):.1f}%  worst {max(eq):.1f}%  n={len(eq)}'])
ws.cell(ws.max_row, 1).font = Font(bold=True, color='1F3864')
ws.append(['"pred (QEMU)" uses no silicon counter for the mix: per-application activity from QEMU '
           'scaled by each application\'s solo iteration rate. "pred (PMU)" re-runs the same model on '
           'silicon-measured activity, isolating model error from emulator error.'])
ws.cell(ws.max_row, 1).font = Font(italic=True, size=9, color='808080')
for col, w in zip('ABCDEFGHIJKLM', (26, 40, 7, 10, 14, 11, 14, 11, 13, 13, 10, 9, 10)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'

# ── bandwidth corner ──────────────────────────────────────────────────────
if os.path.exists('RESULTS-bandwidth.csv'):
    bw = list(csv.DictReader(open('RESULTS-bandwidth.csv')))
    ws = wb.create_sheet('bandwidth corner')
    header(ws, ['case', 'workload', 'GB/s', 'ALU Mops/s', 'pred mW', 'meas mW', 'err %', 'iters/s'])
    for r in sorted(bw, key=lambda r: -float(r['GBps'])):
        e = float(r['err_pct'])
        ws.append([r['case'], r['workload'], float(r['GBps']), float(r['aluMps']),
                   float(r['pred_mW']), float(r['meas_mW']), e, float(r['iters_per_s'])])
        shade(ws.cell(ws.max_row, 7), e if e <= 8 else (9 if e <= 15 else 20))
        ws.cell(ws.max_row, 7).number_format = '0.0'
    ws.append([])
    ws.append(['A pure READ stream at 5.5 GB/s predicts to 1.3%. A 50/50 read+write stream at 11 GB/s '
               'over-predicts 12-13%: the model has ONE bandwidth term, so a written byte is charged the '
               'same as a read byte.'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, color='1F3864')
    ws.append(['Predicted from QEMU activity scaled by the silicon iteration rate; no silicon counter in '
               'the prediction path. Power MEASURED via BCU busy-window median.'])
    ws.cell(ws.max_row, 1).font = Font(italic=True, size=9, color='808080')
    for col, w in zip('ABCDEFGH', (13, 34, 9, 12, 10, 10, 8, 9)):
        ws.column_dimensions[col].width = w

# ── summary ───────────────────────────────────────────────────────────────
ws = wb.create_sheet('summary')
header(ws, ['set', 'rail', 'MAPE %', 'median %', 'p90 %', 'worst %', 'n'])
def agg(vals):
    v = sorted(vals)
    return st.mean(v), st.median(v), v[int(.9 * len(v))], max(v), len(v)
for rail, key in (('vdd_arm', 'err_arm'), ('vdd_soc', 'err_soc'), ('SUM', 'err_sum')):
    m, md, p9, wo, n = agg([float(r[key]) for r in rows])
    ws.append(['50 apps (single)', rail, round(m, 1), round(md, 1), round(p9, 1), round(wo, 1), n])
    shade(ws.cell(ws.max_row, 3), m)
ws.append(['activity: instructions', 'QEMU vs PMU', round(st.mean(ie),2), round(st.median(ie),2),
           round(sorted(ie)[int(.9*len(ie))],2), round(max(ie),2), len(ie)])
shade(ws.cell(ws.max_row, 3), st.mean(ie))
m, md, p9, wo, n = agg(eq)
ws.append(['edge mixes (concurrent)', 'SUM', round(m, 1), round(md, 1), round(p9, 1), round(wo, 1), n])
shade(ws.cell(ws.max_row, 3), m)
allq = [float(r['err_qemu']) for r in mc]
m, md, p9, wo, n = agg(allq)
ws.append(['all mixes (concurrent)', 'SUM', round(m, 1), round(md, 1), round(p9, 1), round(wo, 1), n])
shade(ws.cell(ws.max_row, 3), m)
ws.append([])
ws.append([NOTE]); ws.cell(ws.max_row, 1).font = Font(italic=True, size=9, color='808080')
for col, w in zip('ABCDEFG', (24, 11, 10, 10, 9, 9, 6)):
    ws.column_dimensions[col].width = w

wb.save('RESULTS-50.xlsx')
print('RESULTS-50.xlsx:', wb.sheetnames)
